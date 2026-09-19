"""Build the shareable Word report from the reviewed Markdown source."""

from pathlib import Path
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1] / "reports"
SOURCE = ROOT / "SEPA实验报告_2026-09-16.md"
DESTINATION = SOURCE.with_suffix(".docx")


def set_font(style, size, bold=False, color="243247"):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Microsoft YaHei"
    )


def inline(paragraph, text):
    for part in re.split(r"(\*\*.*?\*\*)", text):
        run = paragraph.add_run(part[2:-2] if part.startswith("**") else part)
        if part.startswith("**"):
            run.bold = True


def field(paragraph, name):
    run = paragraph.add_run()
    element = OxmlElement("w:fldSimple")
    element.set(qn("w:instr"), name)
    run._r.addnext(element)


def table(document, rows):
    headers, *data = rows
    count = len(headers)
    widths = {
        2: [3.5, 13.7],
        3: [5.5, 3.1, 8.6],
        4: [8.2, 3.0, 3.0, 3.0],
        5: [7.2, 2.5, 2.5, 2.5, 2.5],
    }[count]
    if headers[0] == "Encoder 预训练轮数":
        widths = [3.5, 3.2, 3.2, 3.2, 4.1]
    result = document.add_table(rows=1, cols=count)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Cm(width)

    for index, values in enumerate([headers, *data]):
        row = result.rows[0] if index == 0 else result.add_row()
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
        if index == 0:
            props.append(OxmlElement("w:tblHeader"))
        for col, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            fill = OxmlElement("w:shd")
            fill.set(
                qn("w:fill"),
                "213C55" if index == 0 else ("F0F4F7" if index % 2 else "FFFFFF"),
            )
            properties.append(fill)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                item = OxmlElement("w:" + edge)
                item.set(qn("w:w"), "85" if edge in ("top", "bottom") else "95")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            properties.append(margins)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.keep_with_next = index < len(data)
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if col > 0 and count >= 4
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9)
                if index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = Pt(3)
    return result


def main():
    source = SOURCE.read_text(encoding="utf-8")
    document = Document()
    document.core_properties.title = "SEPA 通用视觉表征学习实验报告"
    document.core_properties.subject = "V1、输入视野消融与 V2 阶段总结"
    document.core_properties.author = "SEPA 实验项目"
    document.core_properties.keywords = "SEPA, I-JEPA, ImageNet-100, 实验报告"
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.9)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.7)
    section.header_distance = section.footer_distance = Cm(0.8)

    normal = document.styles["Normal"]
    set_font(normal, 10)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    for name, size in [("Title", 23), ("Heading 1", 15), ("Heading 2", 11.5)]:
        style = document.styles[name]
        set_font(style, size, True, "173750")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(9 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    set_font(document.styles["Subtitle"], 11, color="52697D")
    set_font(document.styles["Caption"], 9, color="52697D")
    document.styles["Caption"].paragraph_format.keep_with_next = True
    document.styles["Caption"].paragraph_format.space_before = Pt(4)
    document.styles["Caption"].paragraph_format.space_after = Pt(4)
    set_font(document.styles["List Bullet"], 10)
    document.styles["List Bullet"].paragraph_format.line_spacing = 1.15
    document.styles["List Bullet"].paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("SEPA 研究记录  /  阶段实验报告")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("718091")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("2026-09-16   ·   ")
    run.font.size = Pt(8)
    field(footer, "PAGE")
    footer.add_run(" / ").font.size = Pt(8)
    field(footer, "NUMPAGES")

    lines = source.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [v.strip() for v in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", v) for v in values):
                    rows.append(values)
                index += 1
            table(document, rows)
            continue
        if line.startswith("# "):
            document.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            heading = document.add_paragraph(line[3:], "Heading 1")
            if line.startswith("## 3."):
                heading.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            document.add_paragraph(line[4:], "Heading 2")
        elif line.startswith("- "):
            inline(document.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"表 \d：", line):
            document.add_paragraph(line, "Caption")
        elif line == "V1、输入视野消融与 V2 阶段总结":
            document.add_paragraph(line, "Subtitle")
        else:
            inline(document.add_paragraph(), line)
        index += 1

    document.save(DESTINATION)
    print(DESTINATION)
    print(f"Tables: {len(document.tables)}; paragraphs: {len(document.paragraphs)}")
    if "--pdf" in sys.argv:
        export_and_preview()


def export_and_preview():
    """Export with a private Word instance, then render a contact sheet for review."""
    command = r"""
$ErrorActionPreference = 'Stop'
$word = $null
$document = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $document = $word.Documents.Open($env:SEPA_REPORT_DOCX, $false, $true)
    $document.Repaginate()
    $pdf = [IO.Path]::ChangeExtension($env:SEPA_REPORT_DOCX, '.pdf')
    $document.ExportAsFixedFormat($pdf, 17)
    Write-Output ('Pages: ' + $document.ComputeStatistics(2))
} finally {
    if ($null -ne $document) { $document.Close(0) }
    if ($null -ne $word) { $word.Quit() }
}
"""
    environment = dict(os.environ)
    environment["SEPA_REPORT_DOCX"] = str(DESTINATION)
    result = subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", command],
        env=environment,
        capture_output=True,
        timeout=120,
    )
    if result.returncode:
        raise RuntimeError(result.stderr.decode(errors="replace"))
    print(result.stdout.decode(errors="replace"))

    import pypdfium2 as pdfium
    from PIL import Image, ImageDraw
    from pypdf import PdfReader

    pdf_path = DESTINATION.with_suffix(".pdf")
    reader = PdfReader(pdf_path)
    visible = "\n".join(page.extract_text() or "" for page in reader.pages)
    for number in ["52.92", "30.34", "46.92", "12,668", "6.02", "2.94"]:
        assert number in visible, number
    print("PDF pages:", len(reader.pages))
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        print("Page", number, "characters:", len(text))

    output = ROOT / ".sepa_report_preview"
    output.mkdir(exist_ok=True)
    pdf = pdfium.PdfDocument(str(pdf_path))
    width, height = 535, 770
    sheet = Image.new("RGB", (width * 3, height * ((len(pdf) + 2) // 3)), "#CBD5DF")
    draw = ImageDraw.Draw(sheet)
    for i in range(len(pdf)):
        rendered = pdf[i].render(scale=1.15).to_pil().convert("RGB")
        rendered.save(output / f"page-{i + 1}.png")
        rendered.thumbnail((width - 20, height - 30))
        x, y = (i % 3) * width + 10, (i // 3) * height + 23
        sheet.paste(rendered, (x, y))
        draw.text((x, y - 17), f"Page {i + 1}", fill="#243247")
    sheet.save(output / "contact-sheet.png")


if __name__ == "__main__":
    main()
